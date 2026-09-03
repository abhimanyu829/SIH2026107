"""PHASE 1 — full analysis of every uploaded file. READ-ONLY on all sources.
Writes only into BIS_SIH26107_DATA/01_ANALYSIS/ ."""
import csv, hashlib, io, json, os, pickle, re, subprocess, sys, datetime
import openpyxl
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from p1_lib import (norm_header, is_empty, is_placeholder, norm_cell, infer_type,
                    multivalue_ratio, semantic_type, score_domains, scheme_col_rank,
                    URL_RE, IS_RE, QCO_RE, DATE_RE)

SRC = "/sessions/optimistic-gallant-gauss/mnt/SIH-DATASETS-EXCEL-SHEET"
UPL = "/sessions/optimistic-gallant-gauss/mnt/uploads"
OUT = os.path.join(SRC, "BIS_SIH26107_DATA", "01_ANALYSIS")
CACHE = os.path.join(OUT, "_cache")
os.makedirs(CACHE, exist_ok=True)
csv.field_size_limit(50_000_000)


def sha256_of(p):
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for b in iter(lambda: f.read(1 << 20), b""):
            h.update(b)
    return h.hexdigest()


def mime_of(p):
    try:
        return subprocess.run(["file", "-b", "--mime-type", p],
                              capture_output=True, text=True, timeout=30).stdout.strip()
    except Exception:
        return "UNKNOWN"


def ts(x):
    return datetime.datetime.utcfromtimestamp(x).strftime("%Y-%m-%dT%H:%M:%SZ")


def enumerate_files():
    files = []
    for root, label in ((SRC, "workspace"), (UPL, "uploads")):
        if not os.path.isdir(root):
            continue
        for name in sorted(os.listdir(root)):
            p = os.path.join(root, name)
            if not os.path.isfile(p):
                continue
            files.append((p, name, label))
    return files


def load_csv(path):
    """Returns (headers, rows, struct) — struct records structural repairs."""
    raw = open(path, "rb").read()
    enc = "utf-8-sig" if raw[:3] == b"\xef\xbb\xbf" else None
    if enc is None:
        try:
            import chardet
            enc = chardet.detect(raw[:200_000]).get("encoding") or "utf-8"
        except Exception:
            enc = "utf-8"
    text = raw.decode(enc, errors="replace")
    struct = {"encoding": enc, "structural_issue": "", "repair_rule": ""}
    # Entire file written as ONE escaped string (literal \n, \") -> repair in memory.
    if "\n" not in text.rstrip("\r\n") and "\\n" in text:
        t = text.strip()
        if t.startswith('"'):
            t = t[1:]
        if t.endswith('"'):
            t = t[:-1]
        text = t.replace('\\"', '"').replace("\\n", "\n")
        struct["structural_issue"] = "SINGLE_LINE_ESCAPED_CSV"
        struct["repair_rule"] = "strip_outer_quotes+unescape(\\\" ,\\n)"
    rows = list(csv.reader(io.StringIO(text)))
    if not rows:
        return [], [], struct
    hdr = [h.replace("﻿", "") for h in rows[0]]
    body = rows[1:]
    width = max([len(hdr)] + [len(r) for r in body]) if body else len(hdr)
    if width > len(hdr):
        struct["structural_issue"] = (struct["structural_issue"] + ";RAGGED_ROWS_WIDER_THAN_HEADER").strip(";")
        hdr = hdr + ["__unnamed_%d" % i for i in range(len(hdr), width)]
    body = [r + [None] * (len(hdr) - len(r)) if len(r) < len(hdr) else r[:len(hdr)] for r in body]
    return hdr, body, struct


def load_xlsx(path):
    """Yields (sheet_name, headers, rows, struct) for EVERY worksheet."""
    wbv = openpyxl.load_workbook(path, data_only=True)
    wbf = openpyxl.load_workbook(path, data_only=False)
    out = []
    for ws in wbv.worksheets:
        wsf = wbf[ws.title]
        formulas = sum(1 for row in wsf.iter_rows() for c in row
                       if isinstance(c.value, str) and c.value.startswith("="))
        links = sum(1 for row in wsf.iter_rows() for c in row if c.hyperlink is not None)
        hidden_r = [str(r) for r, d in ws.row_dimensions.items() if d.hidden]
        hidden_c = [c for c, d in ws.column_dimensions.items() if d.hidden]
        grid = [list(r) for r in ws.iter_rows(values_only=True)]
        hdr = [("" if v is None else str(v)) for v in (grid[0] if grid else [])]
        body = grid[1:]
        while hdr and hdr[-1] == "" and all(is_empty(r[len(hdr) - 1]) for r in body):
            hdr.pop()
            body = [r[:len(hdr)] for r in body]
        struct = {"encoding": "xlsx", "structural_issue": "", "repair_rule": "",
                  "merged_cells": len(ws.merged_cells.ranges),
                  "merged_ranges": ";".join(str(r) for r in list(ws.merged_cells.ranges)[:20]),
                  "formulas": formulas, "hyperlinks": links,
                  "hidden_rows": ";".join(hidden_r[:50]), "hidden_cols": ";".join(hidden_c[:50]),
                  "sheet_state": ws.sheet_state}
        out.append((ws.title, hdr, body, struct))
    wbv.close(); wbf.close()
    return out


# ------------------------------------------------------------------ profiling
def profile_sheet(file_id, fname, sheet, hdr, body):
    n = len(body)
    norm = [norm_header(h) for h in hdr]
    cols, col_rows, scheme_cols = [], [], []
    for j, h in enumerate(hdr):
        raw = [r[j] if j < len(r) else None for r in body]
        vals = [v for v in raw if not is_empty(v)]
        ph = sum(1 for v in vals if is_placeholder(v))
        real = [v for v in vals if not is_placeholder(v)]
        keys = [norm_cell(v) for v in real]
        uniq = len(set(keys))
        dtype = infer_type(real[:4000])
        mv = multivalue_ratio(real[:2000])
        sem = semantic_type(norm[j], dtype, real[:2000], mv)
        # Scheme identity is judged per column, so keep the columns separate.
        if scheme_col_rank(norm[j]):
            scheme_cols.append((norm[j], real[:1500]))
        ex = []
        for v in real:
            s = re.sub(r"\s+", " ", str(v)).strip()
            if s and s not in ex:
                ex.append(s[:70])
            if len(ex) == 3:
                break
        cols.append({"name": h, "norm": norm[j], "dtype": dtype, "sem": sem,
                     "nonnull": len(vals), "ph": ph, "uniq": uniq, "mv": mv})
        col_rows.append([file_id, fname, sheet, j + 1, h, norm[j], dtype, sem,
                         n, len(vals), n - len(vals), ph, uniq,
                         max(0, len(real) - uniq),
                         round(100.0 * len(vals) / n, 2) if n else 0.0,
                         round(100.0 * len(real) / n, 2) if n else 0.0,
                         round(mv, 3), "YES" if mv > 0.30 else "NO",
                         " || ".join(ex)])
    hashes = [hashlib.sha1("\x1f".join(norm_cell(v) for v in r).encode()).hexdigest()[:16]
              for r in body]
    nonblank = sum(1 for r in body if any(not is_empty(v) for v in r))
    nulls = sum(1 for r in body for v in r if is_empty(v))
    phs = sum(c["ph"] for c in cols)
    uniq_rows = len(set(hashes))
    def picks(pred):
        return ";".join(c["name"] for c in cols if pred(c))
    sample = " ".join(str(v) for r in body[:60] for v in r if not is_empty(v))[:20000]
    meta = {
        "row_count": n, "column_count": len(hdr), "headers": " | ".join(hdr),
        "non_empty_rows": nonblank, "non_empty_columns": sum(1 for c in cols if c["nonnull"]),
        "empty_columns": ";".join(c["name"] for c in cols if not c["nonnull"]),
        "null_cell_count": nulls, "placeholder_null_count": phs,
        "unique_row_count": uniq_rows, "duplicate_row_count": n - uniq_rows,
        "data_types": ";".join("%s:%s" % (c["norm"], c["dtype"]) for c in cols),
        "date_columns": picks(lambda c: c["sem"] in ("DATE", "EFFECTIVE_DATE", "YEAR")
                              or c["dtype"] in ("date", "datetime")),
        "url_columns": picks(lambda c: c["sem"] in ("URL", "SOURCE_URL", "DOCUMENT_URL")
                             or c["dtype"] == "url"),
        "is_number_columns": picks(lambda c: c["sem"] in ("IS_NUMBER", "IS_RELATIONSHIP", "STANDARD_ID")),
        "qco_columns": picks(lambda c: c["sem"] in ("QCO_ID", "QCO_REFERENCE")),
        "scheme_columns": picks(lambda c: c["sem"] in ("SCHEME_ID", "SCHEME_NAME", "CERTIFICATION_TYPE")),
        "laboratory_columns": picks(lambda c: c["sem"] in ("LAB_ID", "LAB_IDENTITY", "LAB_CAPABILITY")),
        "multivalue_columns": picks(lambda c: c["mv"] > 0.30),
        "norm_headers": norm, "hashes": hashes, "sample": sample,
        "scheme_cols": scheme_cols,
    }
    return meta, col_rows


def classify(norm_headers, sample, fname, scheme_cols=()):
    sc = score_domains(norm_headers, sample, fname, scheme_cols)
    if not sc:
        return "UNKNOWN", "", 0, "no domain signal in headers or values"
    ranked = sorted(sc.items(), key=lambda kv: -kv[1][0])
    top, (s, ev) = ranked[0]
    if s < 5:
        return "UNKNOWN", ";".join(d for d, _ in ranked[:3]), s, \
               "weak signal (score %d < 5): %s" % (s, ",".join(ev[:4]))
    sec = [d for d, (s2, _) in ranked[1:] if s2 >= 5]
    return top, ";".join(sec[:4]), s, ",".join(ev[:6])


INV_H = ["file_id", "original_filename", "extension", "file_size_bytes", "sha256", "mime_type",
         "created_time_utc", "modified_time_utc", "exact_path", "source_location", "encoding",
         "sheet_count", "total_data_rows", "total_columns", "analysis_status", "candidate_topic",
         "candidate_secondary_topics", "candidate_duplicate_group", "duplicate_group_basis",
         "file_content_signature", "structural_issue", "notes"]
SHEET_H = ["file_id", "original_filename", "sheet_index", "sheet_name", "sheet_state",
           "row_count", "column_count", "headers", "non_empty_rows", "non_empty_columns",
           "empty_columns", "null_cell_count", "placeholder_null_count", "unique_row_count",
           "duplicate_row_count", "merged_cells", "merged_ranges", "formula_cells",
           "hyperlink_cells", "hidden_rows", "hidden_columns", "data_types", "date_columns",
           "url_columns", "is_number_columns", "qco_columns", "scheme_columns",
           "laboratory_columns", "multivalue_columns", "content_sha256_ordered",
           "content_sha256_rowset", "structural_issue", "repair_rule", "notes"]
COL_H = ["file_id", "original_filename", "sheet_name", "column_index", "exact_column_name",
         "normalized_column_name", "data_type", "semantic_type", "total_rows", "non_null_rows",
         "null_rows", "placeholder_null_rows", "unique_count", "duplicate_count",
         "percentage_populated", "percentage_informative", "multivalue_ratio",
         "is_multivalue_field", "example_values"]
CLS_H = ["file_id", "original_filename", "sheet_name", "row_count", "primary_domain",
         "secondary_domains", "confidence_score", "classification_evidence", "notes"]


def main():
    inv, sheets, colrows, clsrows, cache = [], [], [], [], {}
    files = enumerate_files()
    for i, (path, name, loc) in enumerate(files, 1):
        fid = "F%03d" % i
        st = os.stat(path)
        ext = os.path.splitext(name)[1].lower().lstrip(".")
        rec = dict(file_id=fid, original_filename=name, extension=ext,
                   file_size_bytes=st.st_size, sha256=sha256_of(path), mime_type=mime_of(path),
                   created_time_utc=ts(st.st_ctime), modified_time_utc=ts(st.st_mtime),
                   exact_path=path, source_location=loc, encoding="", sheet_count=0,
                   total_data_rows=0, total_columns=0, analysis_status="", candidate_topic="",
                   candidate_secondary_topics="", candidate_duplicate_group="",
                   duplicate_group_basis="", file_content_signature="", structural_issue="",
                   notes="")
        try:
            if ext in ("xlsx", "xlsm", "xltx"):
                loaded = load_xlsx(path)
            elif ext in ("csv", "tsv"):
                h, b, s = load_csv(path)
                loaded = [(os.path.splitext(name)[0], h, b, s)]
            else:
                rec.update(analysis_status="NON_TABULAR_NOT_PROFILED",
                           candidate_topic="OTHER_RELEVANT",
                           notes="Non-tabular file. Inventoried and hashed; no sheet/column "
                                 "profiling applicable. Retained, not deleted.")
                inv.append(rec); print(fid, name, "-> non-tabular"); continue
        except Exception as e:
            rec.update(analysis_status="ERROR_UNREADABLE",
                       notes="load failed: %s: %s" % (type(e).__name__, str(e)[:180]))
            inv.append(rec); print(fid, name, "-> ERROR", e); continue

        rec["sheet_count"] = len(loaded)
        rec["encoding"] = loaded[0][3].get("encoding", "")
        issues, sigs, doms, secs = [], [], {}, set()
        for k, (sname, hdr, body, struct) in enumerate(loaded, 1):
            meta, crows = profile_sheet(fid, name, sname, hdr, body)
            ordered = hashlib.sha256(("|".join(meta["norm_headers"]) + "\n" +
                                     "\n".join(meta["hashes"])).encode()).hexdigest()
            rowset = hashlib.sha256(("|".join(sorted(meta["norm_headers"])) + "\n" +
                                     "\n".join(sorted(meta["hashes"]))).encode()).hexdigest()
            prim, sec, sc, ev = classify(meta["norm_headers"], meta["sample"], name,
                                         meta["scheme_cols"])
            doms[prim] = doms.get(prim, 0) + sc
            secs.update(x for x in sec.split(";") if x)
            if struct.get("structural_issue"):
                issues.append(struct["structural_issue"])
            sigs.append(ordered)
            rec["total_data_rows"] += meta["row_count"]
            rec["total_columns"] += meta["column_count"]
            sheets.append([fid, name, k, sname, struct.get("sheet_state", "visible"),
                           meta["row_count"], meta["column_count"], meta["headers"],
                           meta["non_empty_rows"], meta["non_empty_columns"],
                           meta["empty_columns"], meta["null_cell_count"],
                           meta["placeholder_null_count"], meta["unique_row_count"],
                           meta["duplicate_row_count"], struct.get("merged_cells", 0),
                           struct.get("merged_ranges", ""), struct.get("formulas", 0),
                           struct.get("hyperlinks", 0), struct.get("hidden_rows", ""),
                           struct.get("hidden_cols", ""), meta["data_types"],
                           meta["date_columns"], meta["url_columns"],
                           meta["is_number_columns"], meta["qco_columns"],
                           meta["scheme_columns"], meta["laboratory_columns"],
                           meta["multivalue_columns"], ordered, rowset,
                           struct.get("structural_issue", ""), struct.get("repair_rule", ""),
                           "sheet fully analysed"])
            colrows.extend(crows)
            clsrows.append([fid, name, sname, meta["row_count"], prim, sec, sc, ev,
                            "content-based classification; filename used only as weak signal"])
            cache[(fid, sname)] = {"file": name, "path": path, "headers": hdr,
                                   "norm_headers": meta["norm_headers"],
                                   "hashes": meta["hashes"], "rows": meta["row_count"],
                                   "ordered": ordered, "rowset": rowset,
                                   "primary_domain": prim, "secondary": sec,
                                   "structural_issue": struct.get("structural_issue", "")}
        rec["structural_issue"] = ";".join(sorted(set(issues)))
        rec["analysis_status"] = "ANALYZED_WITH_IN_MEMORY_REPAIR" if issues else "ANALYZED"
        rec["candidate_topic"] = max(doms, key=doms.get) if doms else "UNKNOWN"
        rec["candidate_secondary_topics"] = ";".join(sorted(secs - {rec["candidate_topic"]}))
        rec["file_content_signature"] = hashlib.sha256("|".join(sigs).encode()).hexdigest()
        inv.append(rec)
        print(fid, name, "->", rec["sheet_count"], "sheet(s)", rec["total_data_rows"], "rows",
              rec["candidate_topic"])
    return inv, sheets, colrows, clsrows, cache


def name_family(n):
    s = os.path.splitext(n)[0].lower()
    s = re.sub(r"\s*\(\d+\)\s*$", "", s)
    s = re.sub(r"[_\- ]v\d+$", "", s)
    s = re.sub(r"[_\- ](scheme_\d+|scheme_[ivx]+)$", "", s)
    return re.sub(r"[^a-z0-9]+", "", s)


def assign_groups(inv):
    """Candidate grouping only. Evidence order: identical bytes > identical content > name."""
    by_sha, by_sig, by_fam, g = {}, {}, {}, [0]
    for r in inv:
        by_sha.setdefault(r["sha256"], []).append(r)
        if r["file_content_signature"]:
            by_sig.setdefault(r["file_content_signature"], []).append(r)
        by_fam.setdefault(name_family(r["original_filename"]), []).append(r)
    def tag(groups, basis, prefix):
        for _, rs in sorted(groups.items(), key=lambda kv: kv[1][0]["file_id"]):
            if len(rs) < 2:
                continue
            if any(x["candidate_duplicate_group"] for x in rs):
                continue
            g[0] += 1
            for x in rs:
                x["candidate_duplicate_group"] = "%s-%03d" % (prefix, g[0])
                x["duplicate_group_basis"] = basis
    tag(by_sha, "IDENTICAL_BYTES(sha256)", "EXACTDUP")
    tag(by_sig, "IDENTICAL_TABULAR_CONTENT(all sheets)", "CONTENTDUP")
    tag(by_fam, "FILENAME_FAMILY_ONLY(weak, unconfirmed - Phase 2 must verify by content)",
        "NAMECAND")
    for r in inv:
        if not r["candidate_duplicate_group"]:
            r["candidate_duplicate_group"] = "SINGLETON"
            r["duplicate_group_basis"] = "no byte/content/name peer found in Phase 1"


def write_csv(path, header, rows):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, quoting=csv.QUOTE_MINIMAL)
        w.writerow(header)
        w.writerows(rows)
    return len(rows)


if __name__ == "__main__":
    inv, sheets, colrows, clsrows, cache = main()
    assign_groups(inv)
    n1 = write_csv(os.path.join(OUT, "FILE_INVENTORY.csv"), INV_H,
                   [[r[k] for k in INV_H] for r in inv])
    n2 = write_csv(os.path.join(OUT, "WORKBOOK_SHEET_ANALYSIS.csv"), SHEET_H, sheets)
    n3 = write_csv(os.path.join(OUT, "COLUMN_PROFILING.csv"), COL_H, colrows)
    n4 = write_csv(os.path.join(OUT, "DATASET_CLASSIFICATION.csv"), CLS_H, clsrows)
    with open(os.path.join(CACHE, "phase1_cache.pkl"), "wb") as f:
        pickle.dump({"inventory": inv, "cache": cache}, f)
    summary = {
        "phase": 1, "generated_at_utc": ts(datetime.datetime.utcnow().timestamp()),
        "total_files_enumerated": len(inv),
        "files_analyzed": sum(1 for r in inv if r["analysis_status"].startswith("ANALYZED")),
        "files_repaired_in_memory": sum(1 for r in inv if "REPAIR" in r["analysis_status"]),
        "files_non_tabular": sum(1 for r in inv if r["analysis_status"] == "NON_TABULAR_NOT_PROFILED"),
        "files_error": sum(1 for r in inv if r["analysis_status"] == "ERROR_UNREADABLE"),
        "total_worksheets": len(sheets), "total_columns_profiled": len(colrows),
        "total_data_rows": sum(r["total_data_rows"] for r in inv),
        "classification_rows": len(clsrows),
        "unknown_domain_sheets": sum(1 for r in clsrows if r[4] == "UNKNOWN"),
        "distinct_sha256": len({r["sha256"] for r in inv}),
        "exact_byte_duplicate_files": len(inv) - len({r["sha256"] for r in inv}),
        "rows_written": {"FILE_INVENTORY.csv": n1, "WORKBOOK_SHEET_ANALYSIS.csv": n2,
                         "COLUMN_PROFILING.csv": n3, "DATASET_CLASSIFICATION.csv": n4},
        "source_files_modified": 0,
    }
    with open(os.path.join(OUT, "PHASE1_SUMMARY.json"), "w", encoding="utf-8") as f:
        json.dump(summary, f, indent=2)
    print(json.dumps(summary, indent=2))



