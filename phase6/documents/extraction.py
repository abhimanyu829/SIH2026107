"""Document text extraction: PDF, DOCX, XLSX, TXT (+ image OCR fallback).

Rules (spec section 5):
- text extraction first, always;
- OCR only when extraction is empty/insufficient AND tesseract exists;
- never executes an uploaded file;
- output is page-indexed text so evidence can cite page numbers.
"""
import os
import re


class Page:
    """One page of extracted text (page numbers are 1-based for citations)."""

    __slots__ = ("number", "text", "via")

    def __init__(self, number, text, via="text"):
        self.number, self.text, self.via = number, text or "", via

    def to_dict(self):
        return {"page": self.number, "via": self.via,
                "text": self.text[:12000], "chars": len(self.text)}


def _pages_from_text(text, via="text", page_size=4000):
    """Long plain text is chunked into virtual pages so citations stay useful."""
    text = text or ""
    if not text.strip():
        return []
    if len(text) <= page_size:
        return [Page(1, text, via)]
    pages, i, n = [], 0, 1
    while i < len(text):
        chunk = text[i:i + page_size]
        pages.append(Page(n, chunk, via))
        i += page_size
        n += 1
    return pages


def extract_pdf(path):
    """PyMuPDF text per page; OCR fallback per page only when needed."""
    import pymupdf
    pages = []
    with pymupdf.open(path) as doc:
        for i, page in enumerate(doc, 1):
            text = page.get_text("text") or ""
            if len(text.strip()) < _min_chars(path):
                ocr = _ocr_page(page)
                if ocr:
                    pages.append(Page(i, ocr, via="ocr"))
                    continue
            pages.append(Page(i, text, via="text"))
    return pages


def _min_chars(path):
    import config as cfg
    return cfg.MIN_TEXT_CHARS


def _ocr_page(page):
    """OCR one PDF page image. Returns '' when OCR is unavailable."""
    import config as cfg
    if not cfg.OCR_ENABLED:
        return ""
    try:
        pix = page.get_pixmap(dpi=200)
        img_path = os.path.join(os.environ.get("TEMP", os.getcwd()),
                                "p6_ocr_%d.png" % page.number)
        pix.save(img_path)
        out = _ocr_image_file(img_path)
        try:
            os.remove(img_path)
        except OSError:
            pass
        return out
    except Exception:
        return ""


def _ocr_image_file(img_path):
    """Tesseract OCR of one image file (PNG/JPG uploads + PDF fallback)."""
    import config as cfg
    if not cfg.OCR_ENABLED:
        return ""
    try:
        import pytesseract
        from PIL import Image
        with Image.open(img_path) as im:
            return pytesseract.image_to_string(im) or ""
    except Exception:
        return ""


def extract_docx(path):
    """python-docx: paragraphs + tables (tables often hold test results)."""
    import docx
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return _pages_from_text("\n".join(parts))


def extract_xlsx(path):
    """openpyxl: every sheet as 'Sheet: name' then rows joined with |."""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append("Sheet: %s" % ws.title)
        for row in ws.iter_rows(values_only=True):
            cells = [str(v) for v in row if v is not None and str(v).strip()]
            if cells:
                parts.append(" | ".join(cells))
    wb.close()
    return _pages_from_text("\n".join(parts))


def extract_txt(path):
    with open(path, encoding="utf-8", errors="replace") as f:
        return _pages_from_text(f.read())


def extract_image(path):
    """PNG/JPG: OCR only (these formats carry no text layer)."""
    import config as cfg
    text = _ocr_image_file(path)
    via = "ocr" if text.strip() else "ocr-unavailable"
    return [Page(1, text, via=via)]


EXTRACTORS = {
    ".pdf": extract_pdf,
    ".docx": extract_docx,
    ".xlsx": extract_xlsx,
    ".txt": extract_txt,
    ".png": extract_image,
    ".jpg": extract_image,
    ".jpeg": extract_image,
}


def extract(path):
    """Extract one uploaded file -> (pages, meta). Never raises for bad files;
    returns empty pages + error note instead."""
    meta = {"filename": os.path.basename(path),
            "extension": os.path.splitext(path)[1].lower(),
            "size_bytes": os.path.getsize(path)}
    ext = meta["extension"]
    if ext not in EXTRACTORS:
        return [], dict(meta, error="unsupported file type")
    try:
        pages = EXTRACTORS[ext](path)
    except Exception as e:
        return [], dict(meta, error="%s: %s" % (type(e).__name__, str(e)[:150]))
    meta["pages"] = len(pages)
    meta["chars"] = sum(len(p.text) for p in pages)
    meta["ocr_used"] = any(p.via.startswith("ocr") for p in pages)
    if meta["chars"] < 10:
        meta["error"] = ("no text extracted (OCR unavailable or empty document)"
                         if not meta["ocr_used"] else
                         "OCR produced no usable text")
    return pages, meta
